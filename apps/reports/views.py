from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import urlencode
from xml.sax.saxutils import escape as xml_escape

from django.contrib.auth.decorators import login_required
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date

from apps.accounts.models import UserProfile
from apps.bank_accounts.models import BankAccount
from apps.budgets.models import Budget
from apps.expenses.models import Expense
from apps.income.models import Income


def _get_profile(user):
    profile, _ = UserProfile.objects.get_or_create(
        username=user.username,
        defaults={
            'email': user.email,
            'password_hash': user.password,
            'full_name': user.get_full_name() or user.username,
            'role': UserProfile.ROLE_ADMIN if user.is_staff else UserProfile.ROLE_USER,
            'is_active': user.is_active,
        },
    )
    return profile


def _format_money(amount, currency='VND'):
    amount = Decimal(amount or 0)
    if currency == 'VND':
        return f'{amount:,.0f} {currency}'.replace(',', '.')
    return f'{amount:,.2f} {currency}'


def _money_number(amount):
    return float(Decimal(amount or 0))


def _sum_amount(queryset):
    return queryset.aggregate(total=Sum('amount'))['total'] or Decimal('0')


def _month_bounds(year, month):
    start = date(year, month, 1)
    if month == 12:
        end = date(year + 1, 1, 1)
    else:
        end = date(year, month + 1, 1)
    return start, end


def _selected_period(request):
    today = timezone.now().date()
    year = int(request.GET.get('year') or today.year)
    month = int(request.GET.get('month') or today.month)
    month = min(12, max(1, month))
    return year, month


def _selected_date_range(request):
    today = timezone.now().date()
    default_start = today.replace(day=1)
    start_date = parse_date(request.GET.get('start_date') or '') or default_start
    end_date = parse_date(request.GET.get('end_date') or '') or today
    if start_date > end_date:
        start_date, end_date = end_date, start_date
    return start_date, end_date


def _selected_category_flow(request):
    flow = request.GET.get('flow') or 'expense'
    return flow if flow in {'expense', 'income'} else 'expense'


def _category_flow_label(flow):
    return 'Thu nhập' if flow == 'income' else 'Chi tiêu'


def _monthly_data(profile, year, month):
    currency = profile.default_currency or 'VND'
    start, end = _month_bounds(year, month)
    incomes = Income.objects.filter(
        user=profile,
        status=Income.STATUS_ACTIVE,
        income_date__gte=start,
        income_date__lt=end,
    )
    expenses = Expense.objects.filter(
        user=profile,
        status=Expense.STATUS_ACTIVE,
        expense_date__gte=start,
        expense_date__lt=end,
    )
    income_total = _sum_amount(incomes)
    expense_total = _sum_amount(expenses)
    budget_total = Budget.objects.filter(
        user=profile,
        status=Budget.STATUS_ACTIVE,
        period_month=month,
        period_year=year,
    ).aggregate(total=Sum('spending_limit'))['total'] or Decimal('0')
    balance_total = BankAccount.objects.filter(
        user=profile,
        is_active=True,
    ).aggregate(total=Sum('current_balance'))['total'] or Decimal('0')

    return {
        'currency': currency,
        'income_total': income_total,
        'expense_total': expense_total,
        'budget_total': budget_total,
        'balance_total': balance_total,
        'net_total': income_total - expense_total,
        'incomes': incomes,
        'expenses': expenses,
    }


def _month_keys_between(start_date, end_date):
    current = date(start_date.year, start_date.month, 1)
    end_month = date(end_date.year, end_date.month, 1)
    keys = []

    while current <= end_month and len(keys) < 240:
        keys.append((current.year, current.month))
        if current.month == 12:
            current = date(current.year + 1, 1, 1)
        else:
            current = date(current.year, current.month + 1, 1)

    return keys


def _range_data(profile, start_date, end_date):
    currency = profile.default_currency or 'VND'
    incomes = Income.objects.filter(
        user=profile,
        status=Income.STATUS_ACTIVE,
        income_date__gte=start_date,
        income_date__lte=end_date,
    )
    expenses = Expense.objects.filter(
        user=profile,
        status=Expense.STATUS_ACTIVE,
        expense_date__gte=start_date,
        expense_date__lte=end_date,
    )
    income_total = _sum_amount(incomes)
    expense_total = _sum_amount(expenses)

    budget_filter = Q()
    for year, month in _month_keys_between(start_date, end_date):
        budget_filter |= Q(period_year=year, period_month=month)
    budgets = Budget.objects.filter(user=profile, status=Budget.STATUS_ACTIVE)
    budget_total = budgets.filter(budget_filter).aggregate(total=Sum('spending_limit'))['total'] if budget_filter else Decimal('0')
    budget_total = budget_total or Decimal('0')

    balance_total = BankAccount.objects.filter(
        user=profile,
        is_active=True,
    ).aggregate(total=Sum('current_balance'))['total'] or Decimal('0')

    return {
        'currency': currency,
        'start_date': start_date,
        'end_date': end_date,
        'income_total': income_total,
        'expense_total': expense_total,
        'budget_total': budget_total,
        'balance_total': balance_total,
        'net_total': income_total - expense_total,
        'incomes': incomes,
        'expenses': expenses,
    }


def _category_rows(transactions, currency):
    total = _sum_amount(transactions)
    rows = []
    for row in transactions.values('category__category_name').annotate(total=Sum('amount')).order_by('-total'):
        amount = row['total'] or Decimal('0')
        percent = (amount / total * 100) if total else Decimal('0')
        rows.append(
            {
                'label': row['category__category_name'] or 'Chưa phân loại',
                'value': _format_money(amount, currency),
                'raw_value': _money_number(amount),
                'percent': f'{percent:.1f}'.rstrip('0').rstrip('.'),
            }
        )
    return rows


def _date_range_query_string(start_date, end_date):
    return urlencode(
        {
            'start_date': start_date.isoformat(),
            'end_date': end_date.isoformat(),
        }
    )


def _render_report(request, template_context):
    context = {
        'months': range(1, 13),
        'years': range(timezone.now().year - 4, timezone.now().year + 2),
    }
    context.update(template_context)
    return render(request, 'reports/report_page.html', context)


def _xlsx_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'NUFI Report'
    worksheet.sheet_view.showGridLines = False

    brand_fill = PatternFill('solid', fgColor='1D4ED8')
    dark_fill = PatternFill('solid', fgColor='0F172A')
    section_fill = PatternFill('solid', fgColor='DBEAFE')
    card_fill = PatternFill('solid', fgColor='EFF6FF')
    card_accent_fill = PatternFill('solid', fgColor='CCFBF1')
    header_fill = PatternFill('solid', fgColor='D9F99D')
    zebra_fill = PatternFill('solid', fgColor='F8FAFC')
    border_color = 'CBD5E1'
    thin_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color),
    )
    title_font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    subtitle_font = Font(name='Calibri', size=11, color='E0F2FE')
    section_font = Font(name='Calibri', size=12, bold=True, color='1E3A8A')
    header_font = Font(name='Calibri', size=11, bold=True, color='0F172A')
    body_font = Font(name='Calibri', size=11, color='334155')
    muted_font = Font(name='Calibri', size=10, color='64748B')
    card_label_font = Font(name='Calibri', size=9, bold=True, color='64748B')
    card_value_font = Font(name='Calibri', size=15, bold=True, color='0F172A')
    money_positive_font = Font(name='Calibri', size=11, bold=True, color='15803D')
    money_negative_font = Font(name='Calibri', size=11, bold=True, color='B91C1C')

    def coerce_excel_value(value):
        if isinstance(value, Decimal):
            return float(value), '#,##0.00'
        if hasattr(value, 'strftime'):
            return value, 'dd/mm/yyyy'
        text = str(value or '').strip()
        if text.endswith('VND'):
            number_text = text[:-3].strip().replace(' ', '').replace('+', '').replace('.', '').replace(',', '.')
            try:
                return float(Decimal(number_text)), '#,##0 "VND"'
            except Exception:
                return value, None
        if text.endswith('%'):
            try:
                return float(Decimal(text[:-1].replace(',', '.')) / Decimal('100')), '0.0%'
            except Exception:
                return value, None
        return value, None

    def write_section_title(row_number, title_value):
        worksheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=max_columns)
        cell = worksheet.cell(row=row_number, column=1, value=title_value)
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(horizontal='left', vertical='center')
        worksheet.row_dimensions[row_number].height = 25
        for fill_column in range(2, max_columns + 1):
            worksheet.cell(row=row_number, column=fill_column).fill = section_fill

    def write_summary_cards(start_row, summary_rows):
        if not summary_rows:
            return start_row

        write_section_title(start_row, 'Tổng quan chỉ số')
        card_row = start_row + 2
        for index, summary_row in enumerate(summary_rows):
            label = summary_row[0] if summary_row else ''
            value = summary_row[1] if len(summary_row) > 1 else ''
            column = 1 if index % 2 == 0 else 5
            row_number = card_row + (index // 2) * 4

            worksheet.merge_cells(start_row=row_number, start_column=column, end_row=row_number, end_column=column + 2)
            label_cell = worksheet.cell(row=row_number, column=column, value=label)
            label_cell.fill = card_accent_fill
            label_cell.font = card_label_font
            label_cell.alignment = Alignment(horizontal='left', vertical='center')

            worksheet.merge_cells(start_row=row_number + 1, start_column=column, end_row=row_number + 2, end_column=column + 2)
            excel_value, number_format = coerce_excel_value(value)
            value_cell = worksheet.cell(row=row_number + 1, column=column, value=excel_value)
            value_cell.fill = card_fill
            value_cell.font = card_value_font
            value_cell.alignment = Alignment(horizontal='left', vertical='center')
            if number_format:
                value_cell.number_format = number_format

            for merged_row in range(row_number, row_number + 3):
                for merged_column in range(column, column + 3):
                    cell = worksheet.cell(row=merged_row, column=merged_column)
                    cell.border = thin_border
                    if merged_row > row_number:
                        cell.fill = card_fill

        return card_row + ((len(summary_rows) + 1) // 2) * 4 + 1

    max_columns = max((len(row) for row in rows), default=1)
    max_columns = max(max_columns, 7)

    title_row = rows[0] if rows else ['NUFI Report']
    title = str(title_row[0] or 'NUFI Report')
    subtitle = ' | '.join(str(value) for value in title_row[1:] if value not in (None, ''))

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.fill = brand_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    worksheet.row_dimensions[1].height = 31

    for column_index in range(2, max_columns + 1):
        worksheet.cell(row=1, column=column_index).fill = brand_fill

    current_row = 2
    if subtitle:
        worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_columns)
        subtitle_cell = worksheet.cell(row=2, column=1, value=subtitle)
        subtitle_cell.fill = dark_fill
        subtitle_cell.font = subtitle_font
        subtitle_cell.alignment = Alignment(horizontal='center', vertical='center')
        worksheet.row_dimensions[2].height = 23
        for column_index in range(2, max_columns + 1):
            worksheet.cell(row=2, column=column_index).fill = dark_fill
        current_row = 3

    current_row += 1

    header_keywords = {
        'Chỉ số',
        'Khoản thu/chi',
        'Nội dung',
        'Tháng',
        'Giai đoạn',
        'Danh mục',
        'Mô tả',
        'Mo ta',
        'Chi so',
        'Khoan thu/chi',
    }
    last_header_row = None
    source_rows = list(rows[1:])
    index = 0

    while index < len(source_rows):
        source_row = source_rows[index]
        row_values = list(source_row)
        non_empty_values = [value for value in row_values if value not in (None, '')]

        if not non_empty_values:
            current_row += 1
            index += 1
            continue

        is_section = len(non_empty_values) == 1 and len(row_values) == 1
        is_header = len(row_values) > 1 and str(row_values[0]) in header_keywords
        is_metadata = len(row_values) == 2 and not is_header and not is_section

        if is_header and str(row_values[0]) in {'Chỉ số', 'Chi so'}:
            summary_rows = []
            index += 1
            while index < len(source_rows):
                candidate = list(source_rows[index])
                candidate_values = [value for value in candidate if value not in (None, '')]
                if not candidate_values:
                    break
                if len(candidate_values) == 1 or (len(candidate) > 1 and str(candidate[0]) in header_keywords):
                    break
                summary_rows.append(candidate)
                index += 1
            current_row = write_summary_cards(current_row, summary_rows)
            continue

        if is_section:
            write_section_title(current_row, non_empty_values[0])
            current_row += 1
            index += 1
            continue

        for column_index in range(1, max_columns + 1):
            value = row_values[column_index - 1] if column_index <= len(row_values) else ''
            excel_value, number_format = coerce_excel_value(value)
            cell = worksheet.cell(row=current_row, column=column_index, value=excel_value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if number_format:
                cell.number_format = number_format

            if is_header:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            elif is_metadata and column_index == 1:
                cell.fill = zebra_fill
                cell.font = header_font
            elif is_metadata:
                cell.font = muted_font
            else:
                cell.font = body_font
                if current_row % 2 == 0:
                    cell.fill = zebra_fill

            if hasattr(value, 'strftime'):
                cell.alignment = Alignment(horizontal='center', vertical='center')

            text_value = str(value)
            if 'VND' in text_value or text_value.startswith(('+', '-')):
                if text_value.strip().startswith('-'):
                    cell.font = money_negative_font
                else:
                    cell.font = money_positive_font
                cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)

        if is_header:
            last_header_row = current_row
            worksheet.row_dimensions[current_row].height = 24

        current_row += 1
        index += 1

    worksheet.freeze_panes = 'A4'
    if last_header_row:
        worksheet.auto_filter.ref = f'A{last_header_row}:{get_column_letter(max_columns)}{max(current_row - 1, last_header_row)}'

    for column_index in range(1, max_columns + 1):
        column_letter = get_column_letter(column_index)
        max_length = 0
        for cell in worksheet[column_letter]:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 3, 13), 38)

    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.55
    worksheet.page_margins.bottom = 0.55
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xlsx_response(filename, rows):
    response = HttpResponse(
        _xlsx_bytes(rows),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def overview_summary(request):
    profile = _get_profile(request.user)
    start_date, end_date = _selected_date_range(request)
    data = _range_data(profile, start_date, end_date)
    currency = data['currency']
    period_label = f'{start_date.strftime("%d/%m/%Y")} - {end_date.strftime("%d/%m/%Y")}'
    summary_rows = [
        {'label': 'Tổng tiền vào', 'value': _format_money(data['income_total'], currency), 'note': 'Các khoản thu trong khoảng thời gian đã chọn'},
        {'label': 'Tổng tiền ra', 'value': _format_money(data['expense_total'], currency), 'note': 'Các khoản chi trong khoảng thời gian đã chọn'},
        {'label': 'Chênh lệch', 'value': _format_money(data['net_total'], currency), 'note': 'Tiền vào trừ tiền ra'},
        {'label': 'Ngân sách liên quan', 'value': _format_money(data['budget_total'], currency), 'note': 'Ngân sách của các tháng nằm trong khoảng lọc'},
        {'label': 'Số tiền hiện có', 'value': _format_money(data['balance_total'], currency), 'note': 'Tổng số dư tài khoản đang dùng'},
    ]
    chart_data = {
        'type': 'bar',
        'labels': ['Tiền vào', 'Tiền ra', 'Chênh lệch', 'Ngân sách'],
        'datasets': [
            {
                'label': period_label,
                'data': [
                    _money_number(data['income_total']),
                    _money_number(data['expense_total']),
                    _money_number(data['net_total']),
                    _money_number(data['budget_total']),
                ],
                'backgroundColor': ['#0f766e', '#f97316', '#38bdf8', '#facc15'],
            }
        ],
    }
    return _render_report(
        request,
        {
            'active_report': 'overview',
            'page_title': 'Báo cáo tổng quan',
            'page_intro': 'Tự chọn khoảng thời gian để xem và xuất báo cáo tài chính tổng hợp.',
            'selected_start_date': start_date,
            'selected_end_date': end_date,
            'show_date_range_filter': True,
            'summary_rows': summary_rows,
            'chart_data': chart_data,
            'table_title': f'Tổng quan từ {period_label}',
            'export_excel_url_name': 'export_overview_excel',
            'export_pdf_url_name': 'export_overview_pdf',
            'export_query_string': _date_range_query_string(start_date, end_date),
        },
    )


@login_required
def monthly_summary(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    data = _monthly_data(profile, year, month)
    currency = data['currency']
    summary_rows = [
        {'label': 'Tổng tiền vào', 'value': _format_money(data['income_total'], currency)},
        {'label': 'Tổng tiền ra', 'value': _format_money(data['expense_total'], currency)},
        {'label': 'Chênh lệch', 'value': _format_money(data['net_total'], currency)},
        {'label': 'Ngân sách đã đặt', 'value': _format_money(data['budget_total'], currency)},
        {'label': 'Số tiền hiện có', 'value': _format_money(data['balance_total'], currency)},
    ]
    chart_data = {
        'type': 'bar',
        'labels': ['Tiền vào', 'Tiền ra', 'Ngân sách'],
        'datasets': [
            {
                'label': f'Tháng {month}/{year}',
                'data': [
                    _money_number(data['income_total']),
                    _money_number(data['expense_total']),
                    _money_number(data['budget_total']),
                ],
                'backgroundColor': ['#0f766e', '#f97316', '#facc15'],
            }
        ],
    }
    return _render_report(
        request,
        {
            'active_report': 'monthly',
            'page_title': 'Báo cáo tháng',
            'page_intro': 'Tổng hợp tiền vào, tiền ra, ngân sách và số dư theo tháng.',
            'selected_year': year,
            'selected_month': month,
            'show_month_filter': True,
            'summary_rows': summary_rows,
            'chart_data': chart_data,
            'table_title': f'Tóm tắt tháng {month}/{year}',
            'export_excel_url_name': 'export_monthly_excel',
            'export_pdf_url_name': 'export_monthly_pdf',
        },
    )


@login_required
def yearly_summary(request):
    profile = _get_profile(request.user)
    year, _ = _selected_period(request)
    currency = profile.default_currency or 'VND'
    labels = [f'T{month}' for month in range(1, 13)]
    income_values = []
    expense_values = []
    summary_rows = []

    for month in range(1, 13):
        data = _monthly_data(profile, year, month)
        income_values.append(_money_number(data['income_total']))
        expense_values.append(_money_number(data['expense_total']))
        summary_rows.append(
            {
                'label': f'Tháng {month}',
                'value': _format_money(data['income_total'] - data['expense_total'], currency),
                'note': f"Vào {_format_money(data['income_total'], currency)} | Ra {_format_money(data['expense_total'], currency)}",
            }
        )

    chart_data = {
        'type': 'line',
        'labels': labels,
        'datasets': [
            {
                'label': 'Tiền vào',
                'data': income_values,
                'borderColor': '#0f766e',
                'backgroundColor': 'rgba(15, 118, 110, 0.14)',
                'tension': 0.35,
            },
            {
                'label': 'Tiền ra',
                'data': expense_values,
                'borderColor': '#f97316',
                'backgroundColor': 'rgba(249, 115, 22, 0.14)',
                'tension': 0.35,
            },
        ],
    }
    return _render_report(
        request,
        {
            'active_report': 'yearly',
            'page_title': 'Báo cáo năm',
            'page_intro': 'So sánh tiền vào và tiền ra theo từng tháng trong năm.',
            'selected_year': year,
            'show_month_filter': False,
            'summary_rows': summary_rows,
            'chart_data': chart_data,
            'table_title': f'Tóm tắt năm {year}',
            'export_excel_url_name': 'export_yearly_excel',
            'export_pdf_url_name': 'export_yearly_pdf',
        },
    )


@login_required
def category_expenses(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    flow = _selected_category_flow(request)
    flow_label = _category_flow_label(flow)
    data = _monthly_data(profile, year, month)
    transactions = data['incomes'] if flow == 'income' else data['expenses']
    rows = _category_rows(transactions, data['currency'])
    total_label = 'tổng thu' if flow == 'income' else 'tổng chi'
    chart_data = {
        'type': 'doughnut',
        'labels': [row['label'] for row in rows],
        'datasets': [
            {
                'label': f'{flow_label} theo danh mục',
                'data': [row['raw_value'] for row in rows],
                'backgroundColor': (
                    ['#0f766e', '#22c55e', '#38bdf8', '#14b8a6', '#84cc16', '#06b6d4', '#65a30d']
                    if flow == 'income'
                    else ['#0f766e', '#f59e0b', '#38bdf8', '#f97316', '#a78bfa', '#22c55e', '#f43f5e']
                ),
            }
        ],
    }
    return _render_report(
        request,
        {
            'active_report': 'categories',
            'page_title': f'{flow_label} theo danh mục',
            'page_intro': 'Chọn thu nhập hoặc chi tiêu để xem từng nhóm danh mục đang chiếm tỷ trọng bao nhiêu.',
            'selected_year': year,
            'selected_month': month,
            'selected_flow': flow,
            'flow_options': [
                {'value': 'expense', 'label': 'Chi tiêu'},
                {'value': 'income', 'label': 'Thu nhập'},
            ],
            'show_month_filter': True,
            'summary_rows': [
                {'label': row['label'], 'value': row['value'], 'note': f"{row['percent']}% {total_label}"}
                for row in rows
            ],
            'chart_data': chart_data,
            'table_title': f'Danh mục {flow_label.lower()} tháng {month}/{year}',
            'export_excel_url_name': 'export_categories_excel',
            'export_pdf_url_name': 'export_categories_pdf',
        },
    )


@login_required
def trend_report(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    currency = profile.default_currency or 'VND'
    periods = []
    current_year, current_month = year, month
    for _ in range(6):
        periods.insert(0, (current_year, current_month))
        current_month -= 1
        if current_month == 0:
            current_month = 12
            current_year -= 1

    labels = []
    net_values = []
    summary_rows = []
    for item_year, item_month in periods:
        data = _monthly_data(profile, item_year, item_month)
        labels.append(f'T{item_month}/{item_year}')
        net_values.append(_money_number(data['net_total']))
        summary_rows.append(
            {
                'label': f'Tháng {item_month}/{item_year}',
                'value': _format_money(data['net_total'], currency),
                'note': f"Vào {_format_money(data['income_total'], currency)} | Ra {_format_money(data['expense_total'], currency)}",
            }
        )

    chart_data = {
        'type': 'line',
        'labels': labels,
        'datasets': [
            {
                'label': 'Chênh lệch thu - chi',
                'data': net_values,
                'borderColor': '#0f766e',
                'backgroundColor': 'rgba(15, 118, 110, 0.16)',
                'fill': True,
                'tension': 0.35,
            }
        ],
    }
    return _render_report(
        request,
        {
            'active_report': 'trends',
            'page_title': 'Xu hướng gần đây',
            'page_intro': 'Theo dõi chênh lệch thu - chi trong 6 tháng gần nhất.',
            'selected_year': year,
            'selected_month': month,
            'show_month_filter': True,
            'summary_rows': summary_rows,
            'chart_data': chart_data,
            'table_title': 'Dòng tiền 6 tháng',
            'export_excel_url_name': 'export_trends_excel',
            'export_pdf_url_name': 'export_trends_pdf',
        },
    )


@login_required
def export_monthly_excel(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    data = _monthly_data(profile, year, month)
    currency = data['currency']
    rows = [
        ['Bao cao thang', f'{month}/{year}'],
        ['Chi so', 'Gia tri'],
        ['Tong tien vao', _format_money(data['income_total'], currency)],
        ['Tong tien ra', _format_money(data['expense_total'], currency)],
        ['Chenh lech', _format_money(data['net_total'], currency)],
        ['Ngan sach da dat', _format_money(data['budget_total'], currency)],
        [],
        ['Khoan thu/chi', 'Danh muc', 'So tien', 'Ngay'],
    ]
    for income in data['incomes'].select_related('category'):
        rows.append([income.title, income.category.category_name, _format_money(income.amount, currency), income.income_date])
    for expense in data['expenses'].select_related('category'):
        rows.append([expense.description or 'Khoan chi', expense.category.category_name, _format_money(expense.amount, currency), expense.expense_date])
    return _xlsx_response(f'nufi-{profile.username}-{year}-{month:02d}.xlsx', rows)


@login_required
def export_overview_excel(request):
    profile = _get_profile(request.user)
    start_date, end_date = _selected_date_range(request)
    data = _range_data(profile, start_date, end_date)
    currency = data['currency']
    period_label = f'{start_date.strftime("%d/%m/%Y")} - {end_date.strftime("%d/%m/%Y")}'
    rows = [
        ['Báo cáo tổng quan', period_label],
        ['Người dùng', profile.username],
        [],
        ['Chỉ số', 'Giá trị'],
        ['Tổng tiền vào', _format_money(data['income_total'], currency)],
        ['Tổng tiền ra', _format_money(data['expense_total'], currency)],
        ['Chênh lệch', _format_money(data['net_total'], currency)],
        ['Ngân sách liên quan', _format_money(data['budget_total'], currency)],
        ['Số tiền hiện có', _format_money(data['balance_total'], currency)],
        [],
        ['Chi tiết tiền vào'],
        ['Nội dung', 'Danh mục', 'Tài khoản', 'Số tiền', 'Ngày', 'Ghi chú'],
    ]
    for income in data['incomes'].select_related('category', 'bank_account').order_by('-income_date', '-income_id'):
        rows.append(
            [
                income.title,
                income.category.category_name,
                income.bank_account.account_name if income.bank_account else '',
                _format_money(income.amount, currency),
                income.income_date,
                income.note or '',
            ]
        )

    rows.extend(
        [
            [],
            ['Chi tiết tiền ra'],
            ['Nội dung', 'Danh mục', 'Tài khoản', 'Số tiền', 'Ngày', 'Ghi chú'],
        ]
    )
    for expense in data['expenses'].select_related('category', 'bank_account').order_by('-expense_date', '-expense_id'):
        rows.append(
            [
                expense.description or 'Khoản chi',
                expense.category.category_name,
                expense.bank_account.account_name if expense.bank_account else '',
                _format_money(expense.amount, currency),
                expense.expense_date,
                expense.note or '',
            ]
        )
    return _xlsx_response(
        f'nufi-{profile.username}-overview-{start_date.isoformat()}-{end_date.isoformat()}.xlsx',
        rows,
    )


@login_required
def export_yearly_excel(request):
    profile = _get_profile(request.user)
    year, _ = _selected_period(request)
    currency = profile.default_currency or 'VND'
    rows = [
        ['Báo cáo năm', year],
        ['Người dùng', profile.username],
        [],
        ['Tháng', 'Tiền vào', 'Tiền ra', 'Chênh lệch'],
    ]
    year_income_total = Decimal('0')
    year_expense_total = Decimal('0')
    for month in range(1, 13):
        data = _monthly_data(profile, year, month)
        year_income_total += data['income_total']
        year_expense_total += data['expense_total']
        rows.append(
            [
                f'Tháng {month}',
                _format_money(data['income_total'], currency),
                _format_money(data['expense_total'], currency),
                _format_money(data['net_total'], currency),
            ]
        )
    rows.extend(
        [
            [],
            ['Tổng cả năm', _format_money(year_income_total, currency), _format_money(year_expense_total, currency), _format_money(year_income_total - year_expense_total, currency)],
        ]
    )
    return _xlsx_response(f'nufi-{profile.username}-{year}-yearly.xlsx', rows)


@login_required
def export_categories_excel(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    flow = _selected_category_flow(request)
    flow_label = _category_flow_label(flow)
    data = _monthly_data(profile, year, month)
    currency = data['currency']
    transactions = data['incomes'] if flow == 'income' else data['expenses']
    category_rows = _category_rows(transactions, currency)
    rows = [
        [f'Báo cáo {flow_label.lower()} theo danh mục', f'{month}/{year}'],
        ['Người dùng', profile.username],
        [],
        ['Danh mục', 'Số tiền', 'Tỷ lệ'],
    ]
    for row in category_rows:
        rows.append([row['label'], row['value'], f"{row['percent']}%"])

    rows.extend(
        [
            [],
            [f'Chi tiết {flow_label.lower()}'],
            ['Nội dung', 'Danh mục', 'Tài khoản', 'Số tiền', 'Ngày', 'Ghi chú'],
        ]
    )
    if flow == 'income':
        for income in data['incomes'].select_related('category', 'bank_account').order_by('-income_date', '-income_id'):
            rows.append(
                [
                    income.title,
                    income.category.category_name,
                    income.bank_account.account_name if income.bank_account else '',
                    _format_money(income.amount, currency),
                    income.income_date,
                    income.note or '',
                ]
            )
    else:
        for expense in data['expenses'].select_related('category', 'bank_account').order_by('-expense_date', '-expense_id'):
            rows.append(
                [
                    expense.description or 'Khoản chi',
                    expense.category.category_name,
                    expense.bank_account.account_name if expense.bank_account else '',
                    _format_money(expense.amount, currency),
                    expense.expense_date,
                    expense.note or '',
                ]
            )
    return _xlsx_response(f'nufi-{profile.username}-{year}-{month:02d}-{flow}-categories.xlsx', rows)


@login_required
def export_trends_excel(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    currency = profile.default_currency or 'VND'
    periods = []
    current_year, current_month = year, month
    for _ in range(6):
        periods.insert(0, (current_year, current_month))
        current_month -= 1
        if current_month == 0:
            current_month = 12
            current_year -= 1

    rows = [
        ['Báo cáo xu hướng 6 tháng', f'Kết thúc tháng {month}/{year}'],
        ['Người dùng', profile.username],
        [],
        ['Giai đoạn', 'Tiền vào', 'Tiền ra', 'Chênh lệch'],
    ]
    total_income = Decimal('0')
    total_expense = Decimal('0')
    for item_year, item_month in periods:
        data = _monthly_data(profile, item_year, item_month)
        total_income += data['income_total']
        total_expense += data['expense_total']
        rows.append(
            [
                f'Tháng {item_month}/{item_year}',
                _format_money(data['income_total'], currency),
                _format_money(data['expense_total'], currency),
                _format_money(data['net_total'], currency),
            ]
        )
    rows.extend(
        [
            [],
            ['Tổng 6 tháng', _format_money(total_income, currency), _format_money(total_expense, currency), _format_money(total_income - total_expense, currency)],
        ]
    )
    return _xlsx_response(f'nufi-{profile.username}-{year}-{month:02d}-trends.xlsx', rows)


@login_required
def export_expense_history_excel(request):
    profile = _get_profile(request.user)
    currency = profile.default_currency or 'VND'
    expenses = Expense.objects.filter(user=profile, status=Expense.STATUS_ACTIVE).select_related('category', 'bank_account')

    start_date = request.GET.get('start_date') or ''
    end_date = request.GET.get('end_date') or ''
    category_id = request.GET.get('category') or ''
    bank_account_id = request.GET.get('bank_account') or ''
    if start_date:
        expenses = expenses.filter(expense_date__gte=start_date)
    if end_date:
        expenses = expenses.filter(expense_date__lte=end_date)
    if category_id:
        expenses = expenses.filter(category_id=category_id)
    if bank_account_id:
        expenses = expenses.filter(bank_account_id=bank_account_id)

    rows = [
        ['Lich su chi tieu', profile.username],
        ['Mo ta', 'Danh muc', 'Tai khoan', 'So tien', 'Ngay chi', 'Ghi chu'],
    ]
    for expense in expenses.order_by('-expense_date', '-expense_id'):
        rows.append(
            [
                expense.description or 'Khoan chi',
                expense.category.category_name,
                expense.bank_account.account_name if expense.bank_account else '',
                _format_money(expense.amount, currency),
                expense.expense_date,
                expense.note or '',
            ]
        )
    return _xlsx_response(f'nufi-{profile.username}-expense-history.xlsx', rows)


def _pdf_escape(value):
    return (
        str(value)
        .encode('latin-1', errors='replace')
        .decode('latin-1')
        .replace('\\', '\\\\')
        .replace('(', '\\(')
        .replace(')', '\\)')
    )


def _pdf_color(color):
    return ' '.join(f'{part:.3f}' for part in color)


def _pdf_rect(commands, x, y, width, height, fill, stroke=None):
    commands.append('q')
    commands.append(f'{_pdf_color(fill)} rg')
    if stroke:
        commands.append(f'{_pdf_color(stroke)} RG')
        commands.append('0.8 w')
        commands.append(f'{x:.1f} {y:.1f} {width:.1f} {height:.1f} re B')
    else:
        commands.append(f'{x:.1f} {y:.1f} {width:.1f} {height:.1f} re f')
    commands.append('Q')


def _pdf_line(commands, x1, y1, x2, y2, color, width=0.8):
    commands.append('q')
    commands.append(f'{_pdf_color(color)} RG')
    commands.append(f'{width:.1f} w')
    commands.append(f'{x1:.1f} {y1:.1f} m {x2:.1f} {y2:.1f} l S')
    commands.append('Q')


def _pdf_text(commands, text, x, y, size=11, color=(0.12, 0.18, 0.26), font='F1'):
    commands.extend(
        [
            'BT',
            f'{_pdf_color(color)} rg',
            f'/{font} {size:.1f} Tf',
            f'{x:.1f} {y:.1f} Td',
            f'({_pdf_escape(text)}) Tj',
            'ET',
        ]
    )


def _pdf_fit(text, limit=34):
    text = str(text)
    return text if len(text) <= limit else f'{text[: limit - 3]}...'


def _pdf_document(commands):
    stream = '\n'.join(commands).encode('latin-1', errors='replace')
    objects = [
        b'1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj',
        b'2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj',
        b'3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R /F2 5 0 R >> >> /Contents 6 0 R >> endobj',
        b'4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj',
        b'5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >> endobj',
        b'6 0 obj << /Length ' + str(len(stream)).encode() + b' >> stream\n' + stream + b'\nendstream endobj',
    ]
    pdf = BytesIO()
    pdf.write(b'%PDF-1.4\n')
    offsets = [0]
    for obj in objects:
        offsets.append(pdf.tell())
        pdf.write(obj + b'\n')
    xref_offset = pdf.tell()
    pdf.write(f'xref\n0 {len(objects) + 1}\n'.encode())
    pdf.write(b'0000000000 65535 f \n')
    for offset in offsets[1:]:
        pdf.write(f'{offset:010d} 00000 n \n'.encode())
    pdf.write(f'trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF'.encode())
    return pdf.getvalue()


def _styled_monthly_pdf(profile, year, month, data, currency):
    try:
        return _reportlab_monthly_pdf(profile, year, month, data, currency)
    except ImportError:
        return _basic_monthly_pdf(profile, year, month, data, currency)


def _find_pdf_font(*names):
    font_dirs = [
        Path('/usr/share/fonts/truetype/dejavu'),
        Path('/usr/share/fonts/truetype/liberation'),
        Path('/Library/Fonts'),
        Path('/System/Library/Fonts/Supplemental'),
        Path('/System/Library/Fonts'),
    ]
    for font_dir in font_dirs:
        for name in names:
            font_path = font_dir / name
            if font_path.exists():
                return font_path
    return None


def _register_pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular_path = _find_pdf_font(
        'DejaVuSans.ttf',
        'LiberationSans-Regular.ttf',
        'Arial Unicode.ttf',
        'Arial.ttf',
        'Helvetica.ttc',
    )
    bold_path = _find_pdf_font(
        'DejaVuSans-Bold.ttf',
        'LiberationSans-Bold.ttf',
        'Arial Bold.ttf',
        'Arial Unicode.ttf',
        'Arial.ttf',
    )
    if not regular_path:
        return 'Helvetica', 'Helvetica-Bold'

    regular_name = 'NUFIUnicode'
    bold_name = 'NUFIUnicodeBold'
    registered = set(pdfmetrics.getRegisteredFontNames())
    if regular_name not in registered:
        pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
    if bold_path and bold_name not in registered:
        pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
    return regular_name, bold_name if bold_path else regular_name


def _draw_pdf_text(pdf, text, x, y, font_name, size, color, max_chars=None):
    if max_chars:
        text = _pdf_fit(text, max_chars)
    pdf.setFillColor(color)
    pdf.setFont(font_name, size)
    pdf.drawString(x, y, str(text))


def _reportlab_monthly_pdf(profile, year, month, data, currency):
    category_rows = _category_rows(data['expenses'], currency)
    transaction_rows = []
    for income in data['incomes'].select_related('category').order_by('-income_date', '-income_id'):
        transaction_rows.append(
            [
                income.income_date.strftime('%d/%m/%Y'),
                'Thu nhập',
                income.title,
                income.category.category_name,
                _format_money(income.amount, currency),
            ]
        )
    for expense in data['expenses'].select_related('category').order_by('-expense_date', '-expense_id'):
        transaction_rows.append(
            [
                expense.expense_date.strftime('%d/%m/%Y'),
                'Chi tiêu',
                expense.description or 'Khoản chi',
                expense.category.category_name,
                _format_money(expense.amount, currency),
            ]
        )

    return _reportlab_table_pdf(
        title=f'Báo cáo tài chính tháng {month}/{year}',
        subtitle='Tổng hợp tình hình thu nhập, chi tiêu, ngân sách và số dư trong tháng.',
        profile=profile,
        period_label=f'Tháng {month}/{year}',
        report_code=f'NUFI-MONTHLY-{year}{month:02d}',
        highlights=[
            {'label': 'Tổng tiền vào', 'value': _format_money(data['income_total'], currency), 'note': 'Tổng các khoản thu trong kỳ'},
            {'label': 'Tổng tiền ra', 'value': _format_money(data['expense_total'], currency), 'note': 'Tổng các khoản chi trong kỳ'},
            {'label': 'Chênh lệch', 'value': _format_money(data['net_total'], currency), 'note': 'Tiền vào trừ tiền ra'},
            {'label': 'Ngân sách', 'value': _format_money(data['budget_total'], currency), 'note': 'Tổng hạn mức đã đặt'},
            {'label': 'Số dư hiện tại', 'value': _format_money(data['balance_total'], currency), 'note': 'Tổng số dư tài khoản đang dùng'},
        ],
        sections=[
            {
                'title': 'Bảng tổng hợp tài chính',
                'headers': ['Khoản mục', 'Số tiền', 'Diễn giải'],
                'rows': [
                    ['Tiền vào', _format_money(data['income_total'], currency), 'Tổng thu nhập đã ghi nhận'],
                    ['Tiền ra', _format_money(data['expense_total'], currency), 'Tổng chi tiêu đã ghi nhận'],
                    ['Chênh lệch thu - chi', _format_money(data['net_total'], currency), 'Kết quả dòng tiền trong tháng'],
                    ['Ngân sách đã đặt', _format_money(data['budget_total'], currency), 'Hạn mức theo tháng hoặc danh mục'],
                    ['Số dư hiện tại', _format_money(data['balance_total'], currency), 'Tổng số dư các tài khoản đang hoạt động'],
                ],
                'col_widths': [155, 145, 200],
            },
            {
                'title': 'Chi tiêu theo danh mục',
                'headers': ['Danh mục', 'Số tiền', 'Tỷ trọng'],
                'rows': [[row['label'], row['value'], f"{row['percent']}%"] for row in category_rows],
                'col_widths': [245, 155, 100],
            },
            {
                'title': 'Chi tiết giao dịch trong kỳ',
                'headers': ['Ngày', 'Loại', 'Nội dung', 'Danh mục', 'Số tiền'],
                'rows': transaction_rows,
                'col_widths': [70, 65, 155, 115, 95],
            },
        ],
    )

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    regular_font, bold_font = _register_pdf_fonts()
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    page_width, page_height = A4
    pdf.setTitle(f'Báo cáo tháng {month:02d}/{year} - NUFI')

    brand = colors.HexColor('#0f766e')
    brand_dark = colors.HexColor('#134e4a')
    mint = colors.HexColor('#d9f4ef')
    paper = colors.HexColor('#f5faf7')
    border = colors.HexColor('#d7e3dd')
    text = colors.HexColor('#1f2937')
    muted = colors.HexColor('#6b7280')
    orange = colors.HexColor('#dc6b3f')
    amber = colors.HexColor('#d99a1d')
    blue = colors.HexColor('#357ab8')
    green = colors.HexColor('#288a68')
    white = colors.white

    def rounded_card(x, y, width, height, fill=colors.white, stroke=border, radius=18):
        pdf.setFillColor(fill)
        pdf.setStrokeColor(stroke)
        pdf.setLineWidth(0.7)
        pdf.roundRect(x, y, width, height, radius, fill=1, stroke=1)

    pdf.setFillColor(paper)
    pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)

    pdf.setFillColor(brand)
    pdf.rect(0, page_height - 118, page_width, 118, fill=1, stroke=0)
    pdf.setFillColor(colors.HexColor('#1f9f92'))
    pdf.roundRect(page_width - 205, page_height - 82, 160, 39, 18, fill=1, stroke=0)

    _draw_pdf_text(pdf, 'NUFI', 44, page_height - 49, bold_font, 25, white)
    _draw_pdf_text(pdf, 'Vun đắp tài sản bằng trí tuệ tài chính', 44, page_height - 72, regular_font, 12, colors.HexColor('#dff8f2'))
    _draw_pdf_text(pdf, f'Báo cáo tháng {month:02d}/{year}', 44, page_height - 98, bold_font, 17, white)
    _draw_pdf_text(pdf, f'Người dùng: {_pdf_fit(profile.username, 22)}', page_width - 190, page_height - 61, bold_font, 10, white)
    _draw_pdf_text(pdf, f'Ngày tạo: {timezone.now().strftime("%d/%m/%Y")}', page_width - 190, page_height - 78, regular_font, 9, colors.HexColor('#e6fffb'))

    cards = [
        ('Thu nhập', data['income_total'], green),
        ('Chi tiêu', data['expense_total'], orange),
        ('Chênh lệch', data['net_total'], blue if data['net_total'] >= 0 else orange),
        ('Ngân sách', data['budget_total'], amber),
        ('Số dư hiện tại', data['balance_total'], brand),
    ]
    card_positions = [(42, 638), (214, 638), (386, 638), (42, 548), (214, 548)]
    for (label, value, accent), (x, y) in zip(cards, card_positions):
        rounded_card(x, y, 156, 70)
        pdf.setFillColor(accent)
        pdf.roundRect(x, y, 6, 70, 3, fill=1, stroke=0)
        _draw_pdf_text(pdf, label, x + 18, y + 46, bold_font, 10, muted)
        _draw_pdf_text(pdf, _format_money(value, currency), x + 18, y + 23, bold_font, 12, text, 22)

    _draw_pdf_text(pdf, 'Tổng quan dòng tiền', 42, 505, bold_font, 15, text)
    max_bar = max(
        Decimal(data['income_total'] or 0),
        Decimal(data['expense_total'] or 0),
        Decimal(data['budget_total'] or 0),
        Decimal('1'),
    )
    bars = [
        ('Thu nhập', data['income_total'], green),
        ('Chi tiêu', data['expense_total'], orange),
        ('Ngân sách', data['budget_total'], amber),
    ]
    for index, (label, value, color) in enumerate(bars):
        y = 470 - index * 40
        width = float((Decimal(value or 0) / max_bar) * Decimal(330))
        _draw_pdf_text(pdf, label, 42, y + 5, bold_font, 10, muted)
        pdf.setFillColor(colors.HexColor('#e7efeb'))
        pdf.roundRect(124, y, 350, 14, 7, fill=1, stroke=0)
        pdf.setFillColor(color)
        pdf.roundRect(124, y, max(5, width), 14, 7, fill=1, stroke=0)
        _draw_pdf_text(pdf, _format_money(value, currency), 486, y + 3, regular_font, 9, text, 16)

    _draw_pdf_text(pdf, 'Chi tiêu theo danh mục', 42, 333, bold_font, 15, text)
    pdf.setFillColor(mint)
    pdf.setStrokeColor(border)
    pdf.roundRect(42, 294, 510, 25, 9, fill=1, stroke=1)
    _draw_pdf_text(pdf, 'Danh mục', 58, 303, bold_font, 10, brand_dark)
    _draw_pdf_text(pdf, 'Số tiền', 320, 303, bold_font, 10, brand_dark)
    _draw_pdf_text(pdf, 'Tỷ lệ', 470, 303, bold_font, 10, brand_dark)

    category_rows = _category_rows(data['expenses'], currency)[:6]
    if not category_rows:
        _draw_pdf_text(pdf, 'Chưa có chi tiêu trong tháng này.', 58, 266, regular_font, 11, muted)
    for index, row in enumerate(category_rows):
        y = 266 - index * 28
        pdf.setFillColor(colors.white if index % 2 == 0 else colors.HexColor('#f8fbf8'))
        pdf.rect(42, y - 7, 510, 25, fill=1, stroke=0)
        _draw_pdf_text(pdf, row['label'], 58, y, regular_font, 10, text, 34)
        _draw_pdf_text(pdf, row['value'], 320, y, regular_font, 10, text, 18)
        _draw_pdf_text(pdf, f"{row['percent']}%", 470, y, regular_font, 10, text)

    pdf.setStrokeColor(border)
    pdf.line(42, 58, 552, 58)
    _draw_pdf_text(pdf, 'Báo cáo được tạo từ NUFI - Nurtured Wealth, Financial Intelligence.', 42, 39, regular_font, 9, muted)
    _draw_pdf_text(pdf, 'Số liệu phụ thuộc vào dữ liệu người dùng đã nhập trong hệ thống.', 42, 25, regular_font, 8, muted)

    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _pdf_paragraph(value, style):
    from reportlab.platypus import Paragraph

    safe_value = xml_escape(str(value or '')).replace('\n', '<br/>')
    return Paragraph(safe_value, style)


def _reportlab_table_pdf(title, subtitle, profile, sections, period_label='', report_code='', highlights=None, notes=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

    regular_font, bold_font = _register_pdf_fonts()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=32,
        title=f'{title} - NUFI',
    )

    brand = colors.HexColor('#0f766e')
    brand_dark = colors.HexColor('#134e4a')
    mint = colors.HexColor('#d9f4ef')
    paper = colors.HexColor('#f5faf7')
    border = colors.HexColor('#d7e3dd')
    text = colors.HexColor('#1f2937')
    muted = colors.HexColor('#6b7280')
    warning = colors.HexColor('#dc6b3f')

    styles = {
        'brand': ParagraphStyle('NUFIBrand', fontName=bold_font, fontSize=24, leading=28, textColor=colors.white),
        'header_right': ParagraphStyle('NUFIHeaderRight', fontName=bold_font, fontSize=10, leading=14, textColor=colors.HexColor('#e6fffb'), alignment=TA_RIGHT),
        'header_small': ParagraphStyle('NUFIHeaderSmall', fontName=regular_font, fontSize=8.5, leading=12, textColor=colors.HexColor('#dff8f2')),
        'document_label': ParagraphStyle('NUFIDocumentLabel', fontName=bold_font, fontSize=10, leading=14, textColor=brand_dark, alignment=TA_CENTER, spaceAfter=4),
        'title': ParagraphStyle('NUFITitle', fontName=bold_font, fontSize=18, leading=24, textColor=text, alignment=TA_CENTER, spaceAfter=5),
        'subtitle': ParagraphStyle('NUFISubtitle', fontName=regular_font, fontSize=9.5, leading=14, textColor=muted, alignment=TA_CENTER, spaceAfter=12),
        'section': ParagraphStyle('NUFISection', fontName=bold_font, fontSize=13, leading=17, textColor=brand_dark, spaceBefore=8, spaceAfter=8),
        'cell': ParagraphStyle('NUFICell', fontName=regular_font, fontSize=8.5, leading=12, textColor=text),
        'cell_bold': ParagraphStyle('NUFICellBold', fontName=bold_font, fontSize=8.5, leading=12, textColor=brand_dark),
        'info_label': ParagraphStyle('NUFIInfoLabel', fontName=bold_font, fontSize=8, leading=11, textColor=muted),
        'info_value': ParagraphStyle('NUFIInfoValue', fontName=regular_font, fontSize=8.5, leading=12, textColor=text),
        'card_label': ParagraphStyle('NUFICardLabel', fontName=bold_font, fontSize=7.5, leading=10, textColor=muted),
        'card_value': ParagraphStyle('NUFICardValue', fontName=bold_font, fontSize=11, leading=14, textColor=brand_dark),
        'card_note': ParagraphStyle('NUFICardNote', fontName=regular_font, fontSize=7.5, leading=10, textColor=muted),
        'note': ParagraphStyle('NUFINote', fontName=regular_font, fontSize=8, leading=12, textColor=muted),
        'signature': ParagraphStyle('NUFISignature', fontName=regular_font, fontSize=8.5, leading=12, textColor=text, alignment=TA_CENTER),
    }

    story = []
    generated_at = timezone.now().strftime('%d/%m/%Y %H:%M')
    currency = getattr(profile, 'default_currency', None) or 'VND'
    owner_name = getattr(profile, 'full_name', None) or profile.username
    report_code = report_code or f'NUFI-RPT-{timezone.now().strftime("%Y%m%d%H%M")}'

    header = Table(
        [
            [
                _pdf_paragraph('NUFI', styles['brand']),
                _pdf_paragraph(
                    'BÁO CÁO TÀI CHÍNH CÁ NHÂN\nPersonal Financial Report',
                    styles['header_right'],
                ),
            ],
            [
                _pdf_paragraph('Vun đắp tài sản bằng trí tuệ tài chính', styles['header_small']),
                _pdf_paragraph(f'Mã báo cáo: {report_code}', styles['header_right']),
            ],
        ],
        colWidths=[250, 250],
    )
    header.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, -1), brand),
                ('BOX', (0, 0), (-1, -1), 0, brand),
                ('LEFTPADDING', (0, 0), (-1, -1), 16),
                ('RIGHTPADDING', (0, 0), (-1, -1), 16),
                ('TOPPADDING', (0, 0), (-1, -1), 12),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ]
        )
    )
    story.extend(
        [
            header,
            Spacer(1, 14),
            _pdf_paragraph('BÁO CÁO TÀI CHÍNH CÁ NHÂN', styles['document_label']),
            _pdf_paragraph(title, styles['title']),
            _pdf_paragraph(subtitle, styles['subtitle']),
        ]
    )

    info_table = Table(
        [
            [
                _pdf_paragraph('Người lập báo cáo', styles['info_label']),
                _pdf_paragraph(owner_name, styles['info_value']),
                _pdf_paragraph('Ngày xuất báo cáo', styles['info_label']),
                _pdf_paragraph(generated_at, styles['info_value']),
            ],
            [
                _pdf_paragraph('Kỳ báo cáo', styles['info_label']),
                _pdf_paragraph(period_label or title, styles['info_value']),
                _pdf_paragraph('Đơn vị tiền tệ', styles['info_label']),
                _pdf_paragraph(currency, styles['info_value']),
            ],
            [
                _pdf_paragraph('Tài khoản hệ thống', styles['info_label']),
                _pdf_paragraph(profile.username, styles['info_value']),
                _pdf_paragraph('Nguồn dữ liệu', styles['info_label']),
                _pdf_paragraph('NUFI database', styles['info_value']),
            ],
        ],
        colWidths=[118, 152, 112, 118],
    )
    info_table.setStyle(
        TableStyle(
            [
                ('BACKGROUND', (0, 0), (-1, -1), colors.white),
                ('BOX', (0, 0), (-1, -1), 0.7, border),
                ('INNERGRID', (0, 0), (-1, -1), 0.35, border),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 7),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([info_table, Spacer(1, 12)])

    if highlights:
        story.append(_pdf_paragraph('I. Chỉ số tài chính chính', styles['section']))
        highlight_colors = [brand, warning, colors.HexColor('#357ab8'), colors.HexColor('#d99a1d'), colors.HexColor('#288a68'), brand_dark]
        cards = []
        for index, item in enumerate(highlights):
            label = item.get('label') if isinstance(item, dict) else item[0]
            value = item.get('value') if isinstance(item, dict) else item[1]
            note = item.get('note', '') if isinstance(item, dict) else (item[2] if len(item) > 2 else '')
            accent = item.get('color') if isinstance(item, dict) else None
            accent = accent or highlight_colors[index % len(highlight_colors)]
            card = Table(
                [
                    [_pdf_paragraph(label, styles['card_label'])],
                    [_pdf_paragraph(value, styles['card_value'])],
                    [_pdf_paragraph(note, styles['card_note'])],
                ],
                colWidths=[156],
            )
            card.setStyle(
                TableStyle(
                    [
                        ('BACKGROUND', (0, 0), (-1, -1), colors.white),
                        ('BOX', (0, 0), (-1, -1), 0.7, border),
                        ('LINEBEFORE', (0, 0), (0, -1), 4, accent),
                        ('LEFTPADDING', (0, 0), (-1, -1), 10),
                        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                        ('TOPPADDING', (0, 0), (-1, -1), 5),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                    ]
                )
            )
            cards.append(card)

        for start in range(0, len(cards), 3):
            card_row = cards[start:start + 3]
            while len(card_row) < 3:
                card_row.append('')
            card_table = Table([card_row], colWidths=[166, 166, 166])
            card_table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.extend([card_table, Spacer(1, 8)])

    for index, section in enumerate(sections, start=2 if highlights else 1):
        story.append(_pdf_paragraph(f'{index}. {section["title"]}', styles['section']))
        table_rows = [[_pdf_paragraph(cell, styles['cell_bold']) for cell in section['headers']]]
        body_rows = section.get('rows') or [['Chưa có dữ liệu']]
        for row in body_rows:
            padded_row = list(row) + [''] * (len(section['headers']) - len(row))
            table_rows.append([_pdf_paragraph(cell, styles['cell']) for cell in padded_row[: len(section['headers'])]])

        table = Table(
            table_rows,
            colWidths=section.get('col_widths'),
            repeatRows=1,
            hAlign='LEFT',
        )
        table.setStyle(
            TableStyle(
                [
                    ('BACKGROUND', (0, 0), (-1, 0), mint),
                    ('TEXTCOLOR', (0, 0), (-1, 0), brand_dark),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, paper]),
                    ('GRID', (0, 0), (-1, -1), 0.45, border),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 8),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                    ('TOPPADDING', (0, 0), (-1, -1), 7),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.extend([table, Spacer(1, 12)])

    notes = notes or [
        'Báo cáo được tạo tự động từ dữ liệu người dùng đã nhập trên NUFI.',
        'Các số liệu trong báo cáo dùng cho mục đích theo dõi tài chính cá nhân và hỗ trợ demo hệ thống.',
    ]
    story.append(_pdf_paragraph('Ghi chú báo cáo', styles['section']))
    for note in notes:
        story.append(_pdf_paragraph(f'• {note}', styles['note']))
    story.append(Spacer(1, 16))

    signature_table = Table(
        [
            [
                _pdf_paragraph('Người lập báo cáo', styles['signature']),
                _pdf_paragraph('Người kiểm tra', styles['signature']),
            ],
            ['', ''],
            [
                _pdf_paragraph(owner_name, styles['signature']),
                _pdf_paragraph('____________________', styles['signature']),
            ],
        ],
        colWidths=[250, 250],
        rowHeights=[18, 42, 18],
    )
    signature_table.setStyle(
        TableStyle(
            [
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(signature_table)

    def draw_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular_font, 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(doc.leftMargin, 18, 'NUFI - Nurtured Wealth, Financial Intelligence')
        canvas.drawRightString(A4[0] - doc.rightMargin, 18, f'Trang {canvas.getPageNumber()}')
        canvas.restoreState()

    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return buffer.getvalue()


def _pdf_response(filename, content):
    response = HttpResponse(content, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _basic_monthly_pdf(profile, year, month, data, currency):
    brand = (0.06, 0.37, 0.33)
    brand_dark = (0.04, 0.24, 0.22)
    mint = (0.88, 0.96, 0.94)
    paper = (0.96, 0.98, 0.96)
    border = (0.82, 0.88, 0.84)
    text = (0.12, 0.18, 0.26)
    muted = (0.42, 0.48, 0.58)
    orange = (0.88, 0.42, 0.24)
    amber = (0.92, 0.66, 0.20)
    blue = (0.19, 0.47, 0.75)
    green = (0.16, 0.55, 0.42)

    commands = []
    _pdf_rect(commands, 0, 0, 595, 842, paper)
    _pdf_rect(commands, 0, 724, 595, 118, brand)
    _pdf_rect(commands, 390, 760, 152, 38, (0.15, 0.55, 0.48))
    _pdf_text(commands, 'NUFI', 44, 794, 25, (1, 1, 1), 'F2')
    _pdf_text(commands, 'Vun dap tai san bang tri tue tai chinh', 44, 772, 12, (0.86, 0.96, 0.93), 'F1')
    _pdf_text(commands, f'Bao cao thang {month:02d}/{year}', 44, 748, 17, (1, 1, 1), 'F2')
    _pdf_text(commands, f'Nguoi dung: {_pdf_fit(profile.username, 22)}', 404, 782, 10, (1, 1, 1), 'F2')
    _pdf_text(commands, f'Ngay tao: {timezone.now().strftime("%d/%m/%Y")}', 404, 766, 9, (0.88, 0.98, 0.95), 'F1')

    cards = [
        ('Thu nhap', data['income_total'], green),
        ('Chi tieu', data['expense_total'], orange),
        ('Chen lech', data['net_total'], blue if data['net_total'] >= 0 else orange),
        ('Ngan sach', data['budget_total'], amber),
        ('So du hien tai', data['balance_total'], brand),
    ]
    card_positions = [(42, 638), (214, 638), (386, 638), (42, 548), (214, 548)]
    for (label, value, accent), (x, y) in zip(cards, card_positions):
        _pdf_rect(commands, x, y, 156, 70, (1, 1, 1), border)
        _pdf_rect(commands, x, y, 5, 70, accent)
        _pdf_text(commands, label, x + 18, y + 45, 10, muted, 'F2')
        _pdf_text(commands, _format_money(value, currency), x + 18, y + 22, 12, text, 'F2')

    _pdf_text(commands, 'Tong quan dong tien', 42, 502, 15, text, 'F2')
    max_bar = max(
        Decimal(data['income_total'] or 0),
        Decimal(data['expense_total'] or 0),
        Decimal(data['budget_total'] or 0),
        Decimal('1'),
    )
    bars = [
        ('Thu nhap', data['income_total'], green),
        ('Chi tieu', data['expense_total'], orange),
        ('Ngan sach', data['budget_total'], amber),
    ]
    bar_y = 468
    for index, (label, value, color) in enumerate(bars):
        y = bar_y - index * 40
        width = float((Decimal(value or 0) / max_bar) * Decimal(330))
        _pdf_text(commands, label, 42, y + 8, 10, muted, 'F2')
        _pdf_rect(commands, 124, y, 350, 14, (0.88, 0.92, 0.90))
        _pdf_rect(commands, 124, y, max(4, width), 14, color)
        _pdf_text(commands, _format_money(value, currency), 486, y + 3, 9, text, 'F1')

    _pdf_text(commands, 'Chi tieu theo danh muc', 42, 330, 15, text, 'F2')
    _pdf_rect(commands, 42, 294, 510, 25, mint, border)
    _pdf_text(commands, 'Danh muc', 58, 302, 10, brand_dark, 'F2')
    _pdf_text(commands, 'So tien', 320, 302, 10, brand_dark, 'F2')
    _pdf_text(commands, 'Ty le', 470, 302, 10, brand_dark, 'F2')

    category_rows = _category_rows(data['expenses'], currency)[:6]
    if not category_rows:
        _pdf_text(commands, 'Chua co chi tieu trong thang nay.', 58, 265, 11, muted, 'F1')
    for index, row in enumerate(category_rows):
        y = 264 - index * 28
        fill = (1, 1, 1) if index % 2 == 0 else (0.97, 0.98, 0.96)
        _pdf_rect(commands, 42, y - 6, 510, 25, fill)
        _pdf_text(commands, _pdf_fit(row['label'], 36), 58, y + 1, 10, text, 'F1')
        _pdf_text(commands, row['value'], 320, y + 1, 10, text, 'F1')
        _pdf_text(commands, f"{row['percent']}%", 470, y + 1, 10, text, 'F1')

    _pdf_line(commands, 42, 58, 552, 58, border, 0.8)
    _pdf_text(commands, 'Bao cao duoc tao tu NUFI - Nurtured Wealth, Financial Intelligence.', 42, 38, 9, muted, 'F1')
    _pdf_text(commands, 'So lieu phu thuoc vao du lieu nguoi dung da nhap trong he thong.', 42, 24, 8, muted, 'F1')
    return _pdf_document(commands)


@login_required
def export_monthly_pdf(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    data = _monthly_data(profile, year, month)
    currency = data['currency']
    response = HttpResponse(_styled_monthly_pdf(profile, year, month, data, currency), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="nufi-{profile.username}-{year}-{month:02d}.pdf"'
    return response


@login_required
def export_overview_pdf(request):
    profile = _get_profile(request.user)
    start_date, end_date = _selected_date_range(request)
    data = _range_data(profile, start_date, end_date)
    currency = data['currency']
    period_label = f'{start_date.strftime("%d/%m/%Y")} - {end_date.strftime("%d/%m/%Y")}'
    income_rows = _category_rows(data['incomes'], currency)
    expense_rows = _category_rows(data['expenses'], currency)
    transaction_rows = []

    for income in data['incomes'].select_related('category', 'bank_account').order_by('-income_date', '-income_id'):
        transaction_rows.append(
            [
                income.income_date.strftime('%d/%m/%Y'),
                'Tiền vào',
                income.title,
                income.category.category_name,
                income.bank_account.account_name if income.bank_account else '',
                _format_money(income.amount, currency),
            ]
        )
    for expense in data['expenses'].select_related('category', 'bank_account').order_by('-expense_date', '-expense_id'):
        transaction_rows.append(
            [
                expense.expense_date.strftime('%d/%m/%Y'),
                'Tiền ra',
                expense.description or 'Khoản chi',
                expense.category.category_name,
                expense.bank_account.account_name if expense.bank_account else '',
                _format_money(expense.amount, currency),
            ]
        )

    content = _reportlab_table_pdf(
        title=f'Báo cáo tổng quan {period_label}',
        subtitle='Tổng hợp tình hình tài chính theo khoảng thời gian người dùng lựa chọn.',
        profile=profile,
        period_label=period_label,
        report_code=f'NUFI-OVERVIEW-{start_date.strftime("%Y%m%d")}-{end_date.strftime("%Y%m%d")}',
        highlights=[
            {'label': 'Tổng tiền vào', 'value': _format_money(data['income_total'], currency), 'note': 'Tổng thu trong kỳ'},
            {'label': 'Tổng tiền ra', 'value': _format_money(data['expense_total'], currency), 'note': 'Tổng chi trong kỳ'},
            {'label': 'Chênh lệch', 'value': _format_money(data['net_total'], currency), 'note': 'Tiền vào trừ tiền ra'},
            {'label': 'Ngân sách', 'value': _format_money(data['budget_total'], currency), 'note': 'Các ngân sách liên quan'},
            {'label': 'Số tiền hiện có', 'value': _format_money(data['balance_total'], currency), 'note': 'Tài khoản đang dùng'},
        ],
        sections=[
            {
                'title': 'Tổng hợp tài chính',
                'headers': ['Chỉ số', 'Giá trị', 'Diễn giải'],
                'rows': [
                    ['Tổng tiền vào', _format_money(data['income_total'], currency), 'Các khoản thu trong kỳ'],
                    ['Tổng tiền ra', _format_money(data['expense_total'], currency), 'Các khoản chi trong kỳ'],
                    ['Chênh lệch thu - chi', _format_money(data['net_total'], currency), 'Kết quả dòng tiền'],
                    ['Ngân sách liên quan', _format_money(data['budget_total'], currency), 'Ngân sách của các tháng trong khoảng lọc'],
                    ['Số tiền hiện có', _format_money(data['balance_total'], currency), 'Tổng số dư tài khoản đang hoạt động'],
                ],
                'col_widths': [155, 145, 200],
            },
            {
                'title': 'Tiền vào theo danh mục',
                'headers': ['Danh mục', 'Số tiền', 'Tỷ trọng'],
                'rows': [[row['label'], row['value'], f"{row['percent']}%"] for row in income_rows],
                'col_widths': [245, 155, 100],
            },
            {
                'title': 'Tiền ra theo danh mục',
                'headers': ['Danh mục', 'Số tiền', 'Tỷ trọng'],
                'rows': [[row['label'], row['value'], f"{row['percent']}%"] for row in expense_rows],
                'col_widths': [245, 155, 100],
            },
            {
                'title': 'Chi tiết giao dịch trong kỳ',
                'headers': ['Ngày', 'Loại', 'Nội dung', 'Danh mục', 'Tài khoản', 'Số tiền'],
                'rows': transaction_rows,
                'col_widths': [58, 58, 110, 85, 89, 100],
            },
        ],
    )
    return _pdf_response(
        f'nufi-{profile.username}-overview-{start_date.isoformat()}-{end_date.isoformat()}.pdf',
        content,
    )


@login_required
def export_yearly_pdf(request):
    profile = _get_profile(request.user)
    year, _ = _selected_period(request)
    currency = profile.default_currency or 'VND'
    rows = []
    year_income_total = Decimal('0')
    year_expense_total = Decimal('0')

    for month in range(1, 13):
        data = _monthly_data(profile, year, month)
        year_income_total += data['income_total']
        year_expense_total += data['expense_total']
        rows.append(
            [
                f'Tháng {month}',
                _format_money(data['income_total'], currency),
                _format_money(data['expense_total'], currency),
                _format_money(data['net_total'], currency),
            ]
        )

    content = _reportlab_table_pdf(
        title=f'Báo cáo năm {year}',
        subtitle='So sánh tiền vào, tiền ra và chênh lệch theo từng tháng trong năm.',
        profile=profile,
        period_label=f'Năm {year}',
        report_code=f'NUFI-YEARLY-{year}',
        highlights=[
            {'label': 'Tổng tiền vào', 'value': _format_money(year_income_total, currency), 'note': 'Tổng thu nhập trong năm'},
            {'label': 'Tổng tiền ra', 'value': _format_money(year_expense_total, currency), 'note': 'Tổng chi tiêu trong năm'},
            {'label': 'Chênh lệch năm', 'value': _format_money(year_income_total - year_expense_total, currency), 'note': 'Kết quả dòng tiền cả năm'},
        ],
        sections=[
            {
                'title': 'Tổng quan năm',
                'headers': ['Chỉ số', 'Giá trị'],
                'rows': [
                    ['Tổng tiền vào', _format_money(year_income_total, currency)],
                    ['Tổng tiền ra', _format_money(year_expense_total, currency)],
                    ['Chênh lệch cả năm', _format_money(year_income_total - year_expense_total, currency)],
                ],
                'col_widths': [170, 330],
            },
            {
                'title': 'Chi tiết từng tháng',
                'headers': ['Tháng', 'Tiền vào', 'Tiền ra', 'Chênh lệch'],
                'rows': rows,
                'col_widths': [95, 135, 135, 135],
            },
        ],
    )
    return _pdf_response(f'nufi-{profile.username}-{year}-yearly.pdf', content)


@login_required
def export_categories_pdf(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    flow = _selected_category_flow(request)
    flow_label = _category_flow_label(flow)
    data = _monthly_data(profile, year, month)
    currency = data['currency']
    transactions = data['incomes'] if flow == 'income' else data['expenses']
    category_rows = _category_rows(transactions, currency)
    top_category = category_rows[0]['label'] if category_rows else 'Chưa có dữ liệu'
    detail_rows = []
    if flow == 'income':
        for income in data['incomes'].select_related('category', 'bank_account').order_by('-income_date', '-income_id'):
            detail_rows.append(
                [
                    income.title,
                    income.category.category_name,
                    income.bank_account.account_name if income.bank_account else '',
                    _format_money(income.amount, currency),
                    income.income_date.strftime('%d/%m/%Y'),
                    income.note or '',
                ]
            )
    else:
        for expense in data['expenses'].select_related('category', 'bank_account').order_by('-expense_date', '-expense_id'):
            detail_rows.append(
                [
                    expense.description or 'Khoản chi',
                    expense.category.category_name,
                    expense.bank_account.account_name if expense.bank_account else '',
                    _format_money(expense.amount, currency),
                    expense.expense_date.strftime('%d/%m/%Y'),
                    expense.note or '',
                ]
            )

    content = _reportlab_table_pdf(
        title=f'Báo cáo {flow_label.lower()} theo danh mục tháng {month}/{year}',
        subtitle='Tổng hợp các nhóm danh mục và tỷ trọng của từng nhóm trong tháng.',
        profile=profile,
        period_label=f'Tháng {month}/{year}',
        report_code=f'NUFI-CATEGORY-{flow.upper()}-{year}{month:02d}',
        highlights=[
            {
                'label': f'Tổng {flow_label.lower()}',
                'value': _format_money(data['income_total'] if flow == 'income' else data['expense_total'], currency),
                'note': 'Tổng tiền trong kỳ',
            },
            {'label': 'Số danh mục', 'value': len(category_rows), 'note': f'Danh mục có phát sinh {flow_label.lower()}'},
            {'label': 'Danh mục lớn nhất', 'value': top_category, 'note': 'Danh mục chiếm tỷ trọng cao nhất'},
        ],
        sections=[
            {
                'title': f'Tổng quan {flow_label.lower()} theo danh mục',
                'headers': ['Chỉ số', 'Giá trị'],
                'rows': [
                    [f'Tổng {flow_label.lower()}', _format_money(data['income_total'] if flow == 'income' else data['expense_total'], currency)],
                    [f'Số danh mục có {flow_label.lower()}', len(category_rows)],
                    ['Danh mục lớn nhất', top_category],
                ],
                'col_widths': [170, 330],
            },
            {
                'title': f'{flow_label} theo danh mục',
                'headers': ['Danh mục', 'Số tiền', 'Tỷ lệ'],
                'rows': [[row['label'], row['value'], f"{row['percent']}%"] for row in category_rows],
                'col_widths': [250, 150, 100],
            },
            {
                'title': f'Chi tiết {flow_label.lower()}',
                'headers': ['Mô tả', 'Danh mục', 'Tài khoản', 'Số tiền', 'Ngày', 'Ghi chú'],
                'rows': detail_rows,
                'col_widths': [115, 95, 90, 80, 60, 60],
            },
        ],
    )
    return _pdf_response(f'nufi-{profile.username}-{year}-{month:02d}-{flow}-categories.pdf', content)


@login_required
def export_trends_pdf(request):
    profile = _get_profile(request.user)
    year, month = _selected_period(request)
    currency = profile.default_currency or 'VND'
    periods = []
    current_year, current_month = year, month
    for _ in range(6):
        periods.insert(0, (current_year, current_month))
        current_month -= 1
        if current_month == 0:
            current_month = 12
            current_year -= 1

    rows = []
    total_income = Decimal('0')
    total_expense = Decimal('0')
    for item_year, item_month in periods:
        data = _monthly_data(profile, item_year, item_month)
        total_income += data['income_total']
        total_expense += data['expense_total']
        rows.append(
            [
                f'Tháng {item_month}/{item_year}',
                _format_money(data['income_total'], currency),
                _format_money(data['expense_total'], currency),
                _format_money(data['net_total'], currency),
            ]
        )

    content = _reportlab_table_pdf(
        title=f'Báo cáo xu hướng 6 tháng đến {month}/{year}',
        subtitle='Theo dõi chênh lệch thu - chi trong 6 tháng gần nhất.',
        profile=profile,
        period_label=f'6 tháng đến {month}/{year}',
        report_code=f'NUFI-TREND-{year}{month:02d}',
        highlights=[
            {'label': 'Tổng tiền vào', 'value': _format_money(total_income, currency), 'note': 'Tổng thu nhập 6 tháng'},
            {'label': 'Tổng tiền ra', 'value': _format_money(total_expense, currency), 'note': 'Tổng chi tiêu 6 tháng'},
            {'label': 'Chênh lệch 6 tháng', 'value': _format_money(total_income - total_expense, currency), 'note': 'Xu hướng dòng tiền gần đây'},
        ],
        sections=[
            {
                'title': 'Tổng quan 6 tháng',
                'headers': ['Chỉ số', 'Giá trị'],
                'rows': [
                    ['Tổng tiền vào', _format_money(total_income, currency)],
                    ['Tổng tiền ra', _format_money(total_expense, currency)],
                    ['Chênh lệch 6 tháng', _format_money(total_income - total_expense, currency)],
                ],
                'col_widths': [170, 330],
            },
            {
                'title': 'Dòng tiền theo tháng',
                'headers': ['Giai đoạn', 'Tiền vào', 'Tiền ra', 'Chênh lệch'],
                'rows': rows,
                'col_widths': [110, 130, 130, 130],
            },
        ],
    )
    return _pdf_response(f'nufi-{profile.username}-{year}-{month:02d}-trends.pdf', content)
